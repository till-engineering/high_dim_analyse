"""Gefundene Gruppen beschreiben - gleich fuer PCA, UMAP und t-SNE.

Wo die Gruppen gefunden wurden, ist hier egal: Profil, Vertreter,
Kreuztabellen und Excel werden aus den Rohdaten gerechnet. Bei UMAP und
t-SNE ist das der einzige Weg zu sehen, was eine Insel auf der Karte
ausmacht - die Kartenachsen selbst sind nicht deutbar.

Erzeugt im Ausgabeordner (<praefix> = pca, umap oder tsne):
  <praefix>_gruppenprofil.png   Heatmap: welche Messgroessen machen jede Gruppe aus
  <praefix>_gruppen.xlsx        Zuordnung jeder Excel-Zeile, Profile, Vertreter,
                                Stabilitaet, Zusammenhang mit vorhandenen Spalten
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd

from . import cluster as hc
from .plotstyle import speichere, zeige_oder_schliesse


def zeige_auswahl(cl: hc.Clustering) -> None:
    """Bei KMeans/GMM: die Kennwerte, nach denen die Gruppenzahl gewaehlt wurde."""
    if cl.auswahl is None:
        return
    print("         Gruppenzahl automatisch gewaehlt:")
    for _, z in cl.auswahl.iterrows():
        wahl = "  <-" if z["Gruppen"] == cl.parameter["k"] else ""
        bic = "" if np.isnan(z["BIC"]) else "  BIC %.0f" % z["BIC"]
        print("           k=%d  Silhouette %.2f%s%s"
              % (z["Gruppen"], z["Silhouette"], bic, wahl))
    if cl.auswahl["Silhouette"].max() < 0.25:
        print("         Achtung: Silhouette ueberall unter 0.25 - die Daten "
              "zerfallen nicht klar in Gruppen, die Aufteilung ist willkuerlich.")
    if cl.parameter["k"] == cl.auswahl["Gruppen"].max():
        # Am Rand des Suchbereichs ist das Optimum nicht gefunden, nur abgeschnitten.
        print("         Achtung: gewaehlt ist die groesste gepruefte Gruppenzahl - "
              "die Kennzahl steigt noch, die Zahl ist nicht bestimmt, nur begrenzt. "
              "Mit --gruppen fest vorgeben.")


def beschreibe(cl: hc.Clustering, Z, ds, t, out: Path, praefix: str, verfahren: str,
               laeufe: int = 30, seed: int = 42, quelle: str = "",
               zeigen: bool = False, zusatz: pd.DataFrame | None = None):
    """Konsole, Profil-Heatmap und Excel fuer eine Gruppierung.

    ``Z`` sind die Koordinaten, auf denen geclustert wurde (fuer Stabilitaet
    und Grenzfaelle). ``zusatz`` haengt weitere Spalten an die Zuordnung an,
    z. B. die PCA-Scores oder die Kartenkoordinaten. Liefert die
    Gruppennamen je Zeile, oder ``None``, wenn es keine Gruppen gibt.
    """
    zeige_auswahl(cl)
    if len(cl.gruppen) < 2:
        print("  Keine Gruppenstruktur gefunden (%d Gruppe(n)). Moeglich: die Daten "
              "sind wirklich ein Kontinuum - oder --min-gruppe kleiner waehlen bzw. "
              "--cluster kmeans probieren." % len(cl.gruppen))
        return None

    namen = cl.namen
    anzahl = pd.Series(namen).value_counts()
    stab = (hc.stabilitaet(Z, cl, laeufe, seed=seed)
            if laeufe > 0 else pd.Series(dtype=float))
    prof = hc.profil(ds.X, namen)
    kenn = hc.kennzeichen(prof)
    sil, abstand = hc.einordnung(Z, cl)
    rolle = hc.rollen(cl, sil, abstand)

    # Vorhandene Beschriftungen gehoeren mit in den Abgleich - gerade eine
    # automatisch gewaehlte Label-Spalte kann die Gruppen schon erklaeren.
    meta = ds.meta.copy()
    if ds.labels is not None:
        meta[ds.label_name] = ds.labels
    kreuz = hc.kreuztabellen(meta, namen)
    zeiten = hc.zeitraeume(meta, namen)

    # ---- Konsole: das Wichtigste auf einen Blick ------------------------- #
    for g in prof.columns:
        z = "\n  %-12s n=%d" % (g, anzahl[g])
        if g in stab.index:
            z += "  Stabilitaet %.2f (%s)" % (stab[g], hc.bewertung(stab[g]))
        print(z)
        hoch, tief = kenn[g]
        if hoch:
            print("    hoch    : " + ", ".join("%s %+.1f sd" % e for e in hoch))
        if tief:
            print("    niedrig : " + ", ".join("%s %+.1f sd" % e for e in tief))
        if g != hc.RAUSCHEN:
            typ = ds.zeilen[(namen == g) & (rolle == "typisch")]
            print("    typisch : Excel-Zeilen " + ", ".join(map(str, typ)))
        for spalte in zeiten.columns[::2]:
            basis = spalte[:-len("_von")]
            print("    %s: %s bis %s"
                  % (basis, zeiten.loc[g, spalte], zeiten.loc[g, basis + "_bis"]))
    if kreuz:
        print("\n  Zusammenhang mit vorhandenen Spalten (Cramer V, 1 = deckungsgleich):")
        for spalte, _, v in kreuz:
            print("    %-24s %.2f" % (spalte, v))
    if (stab < 0.6).any():
        print("\n  Achtung: Gruppen mit Stabilitaet unter 0.6 sind vermutlich "
              "Artefakte des Verfahrens - nicht interpretieren.")
    if verfahren in ("UMAP", "t-SNE") and not stab.empty:
        # Neu geclustert wird nur auf der fertigen Karte; die Karte selbst
        # wird nicht neu gerechnet. Deren Zufall steckt also nicht im Wert.
        print("\n  Hinweis: Die Stabilitaet prueft die Gruppensuche auf dieser einen "
              "%s-Karte, nicht die Karte selbst - sie faellt eher zu guenstig aus. "
              "Gegenprobe: --sweep bzw. run_vergleich.py." % verfahren)

    # ---- Profil-Heatmap -------------------------------------------------- #
    print()
    fig = hc.profil_heatmap(prof, anzahl, stab, t, quelle, verfahren)
    speichere(fig, out / ("%s_gruppenprofil.png" % praefix))
    zeige_oder_schliesse(fig, zeigen)

    # ---- Excel zum Nachschlagen und Abgleichen --------------------------- #
    zuordnung = pd.concat([ds.zeilen, meta.reset_index(drop=True)], axis=1)
    zuordnung["Gruppe"] = namen
    zuordnung["Rolle"] = rolle
    zuordnung["Silhouette"] = sil
    zuordnung["Abstand_Mitte"] = abstand
    teile = [zuordnung, ds.X.reset_index(drop=True)]
    if zusatz is not None:
        teile.append(zusatz.reset_index(drop=True))
    zuordnung = pd.concat(teile, axis=1)
    zuordnung = zuordnung.sort_values(["Gruppe", "Abstand_Mitte"], kind="stable")

    uebersicht = pd.DataFrame({"n": anzahl.reindex(prof.columns)})
    uebersicht["Stabilitaet"] = stab.reindex(prof.columns)
    uebersicht["Bewertung"] = [hc.bewertung(v) for v in uebersicht["Stabilitaet"]]
    uebersicht["hoch"] = [", ".join("%s %+.1f" % e for e in kenn[g][0])
                          for g in prof.columns]
    uebersicht["niedrig"] = [", ".join("%s %+.1f" % e for e in kenn[g][1])
                             for g in prof.columns]
    if not zeiten.empty:
        uebersicht = uebersicht.join(zeiten)

    ziel = out / ("%s_gruppen.xlsx" % praefix)
    with pd.ExcelWriter(ziel, engine="openpyxl") as writer:
        uebersicht.to_excel(writer, sheet_name="Uebersicht", index_label="Gruppe")
        zuordnung.to_excel(writer, sheet_name="Zuordnung", index=False)
        zuordnung[zuordnung["Rolle"] != ""].to_excel(
            writer, sheet_name="Vertreter", index=False)
        daten_nach_gruppe(writer, ds, namen, list(prof.columns))
        prof.to_excel(writer, sheet_name="Profil_sigma", index_label="Merkmal")
        hc.mediane(ds.X, namen).to_excel(writer, sheet_name="Profil_Median",
                                         index_label="Merkmal")
        zeile = 0
        for spalte, tab, v in kreuz:
            pd.DataFrame([["%s  (Cramer V = %.2f)" % (spalte, v)]]).to_excel(
                writer, sheet_name="Kreuztabellen", startrow=zeile,
                index=False, header=False)
            tab.to_excel(writer, sheet_name="Kreuztabellen", startrow=zeile + 1)
            zeile += len(tab) + 4
        if cl.auswahl is not None:
            cl.auswahl.to_excel(writer, sheet_name="Auswahl_Gruppenzahl", index=False)
    print("  gespeichert: %s" % ziel)
    return namen


def daten_nach_gruppe(writer, ds, namen, reihenfolge, blatt="Daten_nach_Gruppe"):
    """Alle Rohdaten, Gruppe fuer Gruppe untereinander - zum direkten Durchsehen.

    Die Werte stehen so da wie in der Eingabedatei (Luecken bleiben leer), mit
    allen Spalten in Originalreihenfolge und der Excel-Zeile vorneweg.
    """
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    roh = pd.concat([ds.zeilen, ds.roh], axis=1)
    zeile = 0
    for g in reihenfolge:
        teil = roh[namen == g]
        pd.DataFrame([["%s  (n=%d)" % (g, len(teil))]]).to_excel(
            writer, sheet_name=blatt, startrow=zeile, index=False, header=False)
        teil.to_excel(writer, sheet_name=blatt, startrow=zeile + 1, index=False)
        blattobj = writer.sheets[blatt]
        blattobj.cell(row=zeile + 1, column=1).font = Font(bold=True, size=12)
        zeile += len(teil) + 4        # Titel + Kopfzeile + Daten + 2 Leerzeilen

    # Spaltenbreite grob nach dem laengsten Kopf, sonst ist alles abgeschnitten.
    for i, spalte in enumerate(roh.columns, start=1):
        blattobj.column_dimensions[get_column_letter(i)].width = max(10, len(str(spalte)) + 2)


def cluster_optionen(p) -> None:
    """Die Optionen zur Gruppensuche - gleich fuer PCA, UMAP und t-SNE."""
    g = p.add_argument_group("Gruppen finden (ohne bekannte Labels)")
    g.add_argument("--cluster", default="aus",
                   choices=["aus", "hdbscan", "kmeans", "gmm"],
                   help="Verfahren zur Gruppensuche; hdbscan braucht keine "
                        "Gruppenzahl und laesst Einzelgaenger ohne Gruppe")
    g.add_argument("--gruppen", type=int, default=None,
                   help="feste Gruppenzahl fuer kmeans/gmm (Standard: automatisch)")
    g.add_argument("--min-gruppe", type=int, default=None,
                   help="hdbscan: kleinste Gruppe (Standard: max(5, n/50))")
    g.add_argument("--stabilitaet", type=int, default=30,
                   help="Teilstichproben fuer die Stabilitaetspruefung (0 = aus)")
